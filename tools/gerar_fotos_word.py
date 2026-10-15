#!/usr/bin/env python3
"""Gera documentos DOCX e PDF diretamente a partir de um conjunto de fotos.

Requer: python-docx, Pillow e ReportLab.
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
import tempfile
import zipfile
from collections import deque
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Iterable, Iterator

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    from PIL import Image, ImageOps, UnidentifiedImageError
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas
except ImportError as exc:
    print(
        "Dependência Python ausente. Execute: py -3 -m pip install -r tools/requirements-fotos-word.txt",
        file=sys.stderr,
    )
    raise SystemExit(2) from exc


SUPPORTED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".tif", ".tiff"}
MAX_PHOTOS = 800
MAX_IMAGE_BYTES = 40 * 1024 * 1024
MAX_TOTAL_BYTES = 1024 * 1024 * 1024
MAX_IMAGE_EDGE = 1800
MAX_IMAGE_HEIGHT_CM = 11.2


def natural_key(value: str) -> list[object]:
    return [int(part) if part.isdigit() else part.casefold()
            for part in re.split(r"(\d+)", value)]


def image_sources(input_dir: Path | None, zip_path: Path | None) -> Iterator[tuple[str, bytes]]:
    """Valida nomes/tamanhos antes de processar e carrega uma imagem por vez."""
    if zip_path is not None:
        try:
            with zipfile.ZipFile(zip_path) as archive:
                entries = [
                    item for item in archive.infolist()
                    if not item.is_dir()
                    and Path(item.filename.replace("\\", "/")).suffix.casefold() in SUPPORTED_EXTENSIONS
                ]
                entries.sort(key=lambda item: natural_key(item.filename.replace("\\", "/").split("/")[-1]))
                if len(entries) > MAX_PHOTOS:
                    raise ValueError(f"O arquivo contém mais de {MAX_PHOTOS} imagens. Divida-o em dois arquivos.")

                total_bytes = 0
                for entry in entries:
                    if entry.file_size > MAX_IMAGE_BYTES:
                        raise ValueError(f"A imagem {Path(entry.filename).name} excede 40 MB.")
                    total_bytes += entry.file_size
                    if total_bytes > MAX_TOTAL_BYTES:
                        raise ValueError("O total das imagens excede 1 GB.")
                for entry in entries:
                    # Lê uma única imagem por vez e nunca extrai caminhos do ZIP no disco.
                    yield Path(entry.filename.replace("\\", "/")).name, archive.read(entry)
        except zipfile.BadZipFile as exc:
            raise ValueError("O arquivo enviado não é um ZIP válido.") from exc
    elif input_dir is not None:
        paths = [
            item for item in input_dir.iterdir()
            if item.is_file() and item.suffix.casefold() in SUPPORTED_EXTENSIONS
        ]
        paths.sort(key=lambda item: natural_key(item.name))
        if len(paths) > MAX_PHOTOS:
            raise ValueError(f"Selecione no máximo {MAX_PHOTOS} fotos por vez.")
        total_bytes = 0
        for path in paths:
            size = path.stat().st_size
            if size > MAX_IMAGE_BYTES:
                raise ValueError(f"A imagem {path.name} excede 40 MB.")
            total_bytes += size
            if total_bytes > MAX_TOTAL_BYTES:
                raise ValueError("O total das imagens excede 1 GB.")
        for path in paths:
            yield path.name, path.read_bytes()

def _normalize_one_image(
    index: int,
    filename: str,
    data: bytes,
    temp_dir: Path,
    width_cm: float,
) -> tuple[Path | None, str | None]:
    """Decodifica, redimensiona, gira e salva uma imagem independente."""
    try:
        with Image.open(io.BytesIO(data)) as original:
            image = ImageOps.exif_transpose(original)
            image.load()
            image.thumbnail(
                (MAX_IMAGE_EDGE, MAX_IMAGE_EDGE),
                Image.Resampling.LANCZOS,
                reducing_gap=3.0,
            )

            width_px, height_px = image.size
            if (
                height_px > width_px
                and width_cm * height_px / max(width_px, 1) > MAX_IMAGE_HEIGHT_CM
            ):
                image = image.transpose(Image.Transpose.ROTATE_270)

            has_alpha = (
                image.mode in ("RGBA", "LA")
                or (image.mode == "P" and "transparency" in image.info)
            )
            if has_alpha:
                image = image.convert("RGBA")
                target = temp_dir / f"{index:04d}.png"
                image.save(target, format="PNG", optimize=False, compress_level=3)
            else:
                image = image.convert("RGB")
                target = temp_dir / f"{index:04d}.jpg"
                image.save(target, format="JPEG", quality=86, optimize=False)
            return target, None
    except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError):
        return None, filename


def normalize_images(
    candidates: Iterable[tuple[str, bytes]],
    temp_dir: Path,
    width_cm: float,
) -> list[Path]:
    """Processa imagens em paralelo com fila limitada para controlar a memória."""
    normalized: list[Path] = []
    invalid: list[str] = []
    workers = max(1, min(4, os.cpu_count() or 1))
    max_pending = workers * 2
    source = iter(enumerate(candidates, start=1))
    pending = deque()

    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="foto") as executor:
        # Mantemos no máximo 2 tarefas por trabalhador em memória.
        for _ in range(max_pending):
            try:
                index, (filename, data) = next(source)
            except StopIteration:
                break
            pending.append(
                executor.submit(_normalize_one_image, index, filename, data, temp_dir, width_cm)
            )

        while pending:
            future = pending.popleft()
            target, filename_invalido = future.result()
            if target is not None:
                normalized.append(target)
            if filename_invalido is not None:
                invalid.append(filename_invalido)

            try:
                index, (filename, data) = next(source)
            except StopIteration:
                continue
            pending.append(
                executor.submit(_normalize_one_image, index, filename, data, temp_dir, width_cm)
            )

    if not normalized:
        raise ValueError("Os arquivos enviados não contêm imagens válidas.")
    if invalid:
        print(
            f"Aviso: {len(invalid)} arquivo(s) inválido(s) ou corrompido(s) foram ignorados.",
            file=sys.stderr,
        )
    return normalized

def image_dimensions(image_path: Path, width_cm: float) -> tuple[float, float]:
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    scale = min(
        width_cm / max(width_px, 1),
        MAX_IMAGE_HEIGHT_CM / max(height_px, 1),
    )
    return width_px * scale, height_px * scale


def build_docx(images: list[Path], output_path: Path, width_cm: float) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)

    # Duas fotos por página; cada imagem cabe em uma caixa com altura máxima.
    for index, image_path in enumerate(images):
        largura_final, altura_final = image_dimensions(image_path, width_cm)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(12)
        paragraph.paragraph_format.keep_together = True
        paragraph.add_run().add_picture(
            str(image_path),
            width=Cm(largura_final),
            height=Cm(altura_final),
        )
        if (index + 1) % 2 == 0 and index + 1 < len(images):
            document.add_page_break()

    document.save(output_path)


def build_pdf(images: list[Path], output_path: Path, width_cm: float) -> None:
    """Renderiza o PDF diretamente, sem abrir/converter o DOCX no LibreOffice."""
    page_width, page_height = A4
    max_height = MAX_IMAGE_HEIGHT_CM * cm
    side_margin = 1.8 * cm
    top_margin = 1.3 * cm
    bottom_margin = 1.3 * cm
    width_limit = min(width_cm * cm, page_width - (2 * side_margin))

    pdf = canvas.Canvas(str(output_path), pagesize=A4, pageCompression=1)
    for index, image_path in enumerate(images):
        draw_width_cm, draw_height_cm = image_dimensions(image_path, width_limit / cm)
        draw_width = draw_width_cm * cm
        draw_height = draw_height_cm * cm

        # Cada página tem duas áreas fixas; centraliza a foto na sua área.
        slot_bottom = (page_height - top_margin - max_height) if index % 2 == 0 else bottom_margin
        draw_x = (page_width - draw_width) / 2
        draw_y = slot_bottom + (max_height - draw_height) / 2
        pdf.drawImage(
            str(image_path),
            draw_x,
            draw_y,
            width=draw_width,
            height=draw_height,
            preserveAspectRatio=True,
            anchor="c",
            mask="auto",
        )

        if (index + 1) % 2 == 0 and index + 1 < len(images):
            pdf.showPage()

    pdf.save()

def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path)
    parser.add_argument("--zip-file", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--width", type=float, default=12.0)
    args = parser.parse_args()

    if (args.input_dir is None) == (args.zip_file is None):
        raise ValueError("Informe uma pasta de imagens ou um arquivo ZIP.")
    if not 5.0 <= args.width <= 17.4:
        raise ValueError("A largura deve estar entre 5 cm e 17,4 cm.")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    docx_path = args.output_dir / "fotos_redimensionadas.docx"
    pdf_path = args.output_dir / "fotos_redimensionadas.pdf"

    candidates = image_sources(args.input_dir, args.zip_file)
    with tempfile.TemporaryDirectory(prefix="fotos-word-images-") as temp:
        normalized = normalize_images(candidates, Path(temp), args.width)
        # Word e PDF são documentos independentes: monte os dois em paralelo.
        with ThreadPoolExecutor(max_workers=2, thread_name_prefix="documento") as executor:
            docx_future = executor.submit(build_docx, normalized, docx_path, args.width)
            pdf_future = executor.submit(build_pdf, normalized, pdf_path, args.width)
            docx_future.result()
            pdf_future.result()
        stats = {
            "ok": True,
            "photos": len(normalized),
            "width_cm": args.width,
            "docx": docx_path.name,
            "pdf": pdf_path.name,
        }
        print(json.dumps(stats, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        raise SystemExit(1)
